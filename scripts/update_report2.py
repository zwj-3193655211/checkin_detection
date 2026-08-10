"""
实训报告综合改进（第二批）：
1. 3.1.1 测试环境 Python 3.8+/PyTorch 1.9+ → Python 3.10/PyTorch 2.x
2. 2.6.5.4 "首次将"绝对化表述弱化
3. 2.4.2 温度选择依据补充特征预测器 T=1.8
4. 3.5.1 消融实验严谨性说明（全集=拟合效果）
5. 3.5.1 新增实验五：11独立MLP对比实验
6. 3.7 独立测试补充严谨性注（未人工复核、无法计算实测漏检率）
7. 3.5.2 阈值分析表述与0.80/0.85双阈值对齐
8. 5.4.1 特征依赖性补充实验印证
9. 5.5 口语化结尾改写
"""
from copy import deepcopy

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

DOC = '实训项目实践报告.docx'


def replace_in_runs(par, old, new):
    full = ''.join(r.text for r in par.runs)
    if old not in full:
        return False
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    runs = list(par.runs)
    first_idx = 0
    for i, r in enumerate(runs):
        if r.text.strip():
            first_idx = i
            break
    runs[first_idx].text = full.replace(old, new)
    for r in runs[:first_idx] + runs[first_idx + 1:]:
        r.text = ''
    return True


def replace_par_exact(par, old, new):
    full = ''.join(r.text for r in par.runs)
    if full.strip() != old.strip():
        return False
    runs = list(par.runs)
    runs[0].text = new
    for r in runs[1:]:
        r.text = ''
    return True


def insert_paragraph_after(paragraph, text):
    """在指定段落后插入新段落（复制原段落格式）"""
    new_p = deepcopy(paragraph._p)
    # 清空内容只保留段落格式
    for child in list(new_p):
        if child.tag != qn('w:pPr'):
            new_p.remove(child)
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.add_run(text)
    return new_para


def main():
    d = docx.Document(DOC)
    stats = {'para': 0, 'insert': 0}

    for p in d.paragraphs:
        full = ''.join(r.text for r in p.runs)

        # 1. 测试环境
        if '软件环境：Python 3.8+，PyTorch 1.9+，OpenAI CLIP' in full:
            replace_in_runs(p, 'Python 3.8+，PyTorch 1.9+', 'Python 3.10，PyTorch 2.x')
            stats['para'] += 1
        # 2. "首次将"弱化（置信度正则化损失）
        elif '首次将Focal Loss与熵正则化结合' in full:
            replace_in_runs(p, '首次将Focal Loss与熵正则化结合', '将Focal Loss与熵正则化结合')
            stats['para'] += 1
        # 3. 温度选择依据补充
        elif '选择依据：通过消融实验，T=5.0 时在准确率和审核率之间取得最佳平衡。' in full:
            replace_in_runs(p, 'T=5.0 时在准确率和审核率之间取得最佳平衡',
                            '主分类器温度T=5.0在准确率和审核率之间取得最佳平衡；特征预测器采用T=1.8平衡特征概率分布')
            stats['para'] += 1
        # 4. 消融实验严谨性说明
        elif '注：消融实验使用全数据集（2067张）以验证各参数的边际影响。' in full:
            replace_in_runs(p, '以验证各参数的边际影响',
                            '（含训练集）以验证各参数的边际影响，该结果反映参数变化对数据集的拟合效果，'
                            '而非泛化能力的绝对评估；各配置在独立测试集上的最终性能对比见3.4节')
            stats['para'] += 1
        # 7. 3.5.2 阈值分析表述
        elif '分析：阈值 0.85 是最优选择，既保证零漏检，又保持较高准确率。' in full:
            replace_in_runs(p, '分析：阈值 0.85 是最优选择，既保证零漏检，又保持较高准确率',
                            '分析：历史消融中单一阈值0.85即能保证零漏检并保持较高准确率；'
                            '当前系统进一步采用0.80/0.85双阈值设计，在保证零漏检的同时降低审核率')
            stats['para'] += 1
        # 8. 特征依赖性补充
        elif '特征依赖性：系统高度依赖特征识别准确性' in full:
            replace_in_runs(p, '系统高度依赖特征识别准确性',
                            '系统高度依赖特征识别准确性（实验证实：独立特征模型误报使自动通过样本中'
                            '错误特征依赖比例升至4.5%，故保留共享结构）')
            stats['para'] += 1
        # 9. 口语化结尾
        elif '此外，在实验过程中，我感受到了选型、设计方案、调参的困扰' in full:
            new_text = ('此外，在实验过程中，我经历了方案选型、模型设计与参数调优的挑战，'
                        '这促使我主动查阅资料、请教老师，最终积累了宝贵的工程实践经验。')
            replace_par_exact(p, full.strip(), new_text)
            stats['para'] += 1

    # 5. 插入实验五（P478 分析段之后 = 实验四分析）
    exp4_anchor = None
    for p in d.paragraphs:
        if '号码布是晨跑的关键特征，增强后得分从0.346提升至0.677' in p.text:
            exp4_anchor = p
            break
    if exp4_anchor is not None:
        blocks = [
            '实验五：11独立MLP特征预测器的影响',
            '设置：将11维共享特征预测器拆分为11个独立小MLP（每特征512→256→128→1），'
            '每特征独立设置Focal Loss权重与温度，与共享结构对比。',
            '结果：①特征准确率与共享结构基本持平（测试集11项中独立模型5项略优、3项持平）；'
            '②自动通过样本中依赖错误特征预测的比例从0.5%升至4.5%，特征误报率全面上升'
            '（如跑道0.31%→3.12%），说明独立模型在少数类特征上更易过度自信；'
            '③三支决策审核率随概率校准方式波动剧烈（保守训练85.1%、激进训练15.9%），异常泛化稳定性低于共享结构。',
            '分析：共享网络可利用晨读/晨跑特征间的语义相关性互相强化，小样本特征（号码布训练集仅56个正样本）'
            '在共享表示下训练更稳定；独立模型在异常样本上的泛化更脆弱（验证集不漏检但测试集出现漏检）。'
            '故最终保留共享结构，特征识别优化方向改为收集更多标注样本。',
        ]
        anchor = exp4_anchor
        for b in blocks:
            anchor = insert_paragraph_after(anchor, b)
        stats['insert'] += 4

    # 6. 3.7 独立测试补充注（"最终测试结果如下："之后）
    p489 = None
    for p in d.paragraphs:
        if '产生了一个共计613张图片未经人工数据标注的数据集' in p.text:
            p489 = p
            break
    if p489 is not None:
        note = ('注：该独立实测仅统计系统的自动通过/待审核决策数量，未对自动通过样本逐一人工复核，'
                '无法直接计算实测漏检率；系统设计上所有低置信度/少特征样本均进入待审核队列，'
                '异常图片被自动放行的风险由规则3/4约束。')
        insert_paragraph_after(p489, note)
        stats['insert'] += 1

    d.save(DOC)
    print(f"完成: 段落修改 {stats['para']} 处, 插入 {stats['insert']} 处")


if __name__ == '__main__':
    main()
